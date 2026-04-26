#!/usr/bin/env python3

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent / "utilcupos-kickoff.pptx"

NAVY = RGBColor(0x1E, 0x29, 0x3A)
TEAL = RGBColor(0x10, 0xC1, 0xBC)
ORANGE = RGBColor(0xFF, 0x91, 0x41)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x56, 0x56, 0x56)
LIGHT = RGBColor(0xF2, 0xF6, 0xF7)


def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, title):
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()

    accent = slide.shapes.add_shape(1, Inches(0), Inches(0.85), Inches(13.333), Inches(0.05))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(12.2), Inches(0.45))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Arial"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE


def add_title_slide(slide, title, subtitle, meta):
    add_bg(slide, NAVY)
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(12), Inches(1.5))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE

    tb2 = slide.shapes.add_textbox(Inches(1.1), Inches(2.6), Inches(11.2), Inches(0.8))
    p = tb2.text_frame.paragraphs[0]
    p.text = subtitle
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.color.rgb = TEAL

    tb3 = slide.shapes.add_textbox(Inches(1.2), Inches(4.0), Inches(11), Inches(1.6))
    tf = tb3.text_frame
    for i, line in enumerate(meta):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.color.rgb = WHITE

    bar2 = slide.shapes.add_shape(1, Inches(0), Inches(7.42), Inches(13.333), Inches(0.08))
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = ORANGE
    bar2.line.fill.background()


def add_bullets_slide(slide, title, bullets, note=None):
    add_bg(slide, WHITE)
    add_header(slide, title)
    box = slide.shapes.add_shape(1, Inches(0.55), Inches(1.2), Inches(12.2), Inches(5.7))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = TEAL

    tb = slide.shapes.add_textbox(Inches(0.9), Inches(1.45), Inches(11.5), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.name = "Arial"
        p.font.size = Pt(20)
        p.font.color.rgb = NAVY
        p.level = 0
        p.space_after = Pt(10)

    if note:
        tb2 = slide.shapes.add_textbox(Inches(0.9), Inches(6.35), Inches(11.5), Inches(0.45))
        p = tb2.text_frame.paragraphs[0]
        p.text = note
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.italic = True
        p.font.color.rgb = GRAY


def add_table_slide(slide, title, headers, rows, widths):
    add_bg(slide, WHITE)
    add_header(slide, title)
    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.45), Inches(1.35), Inches(12.4), Inches(5.8)).table
    for idx, text in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        para = cell.text_frame.paragraphs[0]
        para.font.name = "Arial"
        para.font.size = Pt(13)
        para.font.bold = True
        para.font.color.rgb = WHITE
        table.columns[idx].width = Inches(widths[idx])

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if r_idx % 2 else WHITE
            para = cell.text_frame.paragraphs[0]
            para.font.name = "Arial"
            para.font.size = Pt(12)
            para.font.color.rgb = NAVY


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_slide(
        s,
        "WO0000000358906 - Utilizaciones de Cupos",
        "Kickoff del proyecto",
        [
            "Cliente: Bancoomeva",
            "Inicio oficial: 20-abr-2026",
            "Fecha compromiso: 20-may-2026",
            "Project Manager: Valeria Rivera Rico",
        ],
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bullets_slide(
        s,
        "Objetivo y Alcance",
        [
            "Implementar el modulo de Utilizacion de Cupos para Bancoomeva.",
            "Prioridad inmediata: entregar utilizarCuposOrq y avanzar en paralelo las HU funcionales HU-BCO-004 a HU-BCO-015.",
            "La Epica 1 no hace parte de la prioridad inmediata y se abordara mas adelante.",
        ],
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_table_slide(
        s,
        "Equipo Clave",
        ["Rol", "Nombre", "Dedicacion"],
        [
            ["Gerente de Proyecto", "Valeria Rivera Rico", "4 horas diarias"],
            ["Soporte PM", "Raul Collazos Castillo", "Bajo demanda"],
            ["PO VortexBird", "Ingrid Y. Mosquera", "Bajo demanda"],
            ["PO / Analista funcional cliente", "Adriana Munoz Nanez", "Cliente"],
            ["Arquitecto", "Christian Ospina", "Por definir"],
            ["Apoyo arquitectura", "Sebastian Garcia", "Por definir"],
            ["QA", "Kathereen Gonzalez", "100% tras cierre de desarrollo"],
        ],
        [3.2, 5.2, 3.7],
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_table_slide(
        s,
        "Distribucion de Trabajo",
        ["Frente", "Responsable", "Detalle"],
        [
            ["Orquestacion", "Jurgen Sanclemente", "Servicio utilizarCuposOrq"],
            ["Funcional", "Sebastian Caicedo", "HU-BCO-004 a HU-BCO-015"],
            ["Funcional", "David Velazco", "HU-BCO-004 a HU-BCO-015"],
            ["QA", "Kathereen Gonzalez", "Inicia despues del desarrollo"],
        ],
        [2.5, 3.6, 6.0],
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bullets_slide(
        s,
        "Mensajes Clave del Kickoff",
        [
            "1. Se debe entregar el servicio utilizarCuposOrq sin la mediacion el lunes de la proxima semana.",
            "2. El frente banca debe hacer la solicitud formal de la mediacion del servicio utilizarCupos a la celula de Interoperabilidad Bancoomeva.",
            "3. La mediacion no hace parte del alcance del proyecto.",
            "4. Sebastian Caicedo y David Velazco se comprometieron a entregar su parte el 04-may-2026 para iniciar QA.",
        ],
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_table_slide(
        s,
        "Cronograma Ejecutivo",
        ["Hito", "Fecha", "Observacion"],
        [
            ["Inicio oficial", "20-abr-2026", "Confirmado"],
            ["Entrega utilizarCuposOrq sin mediacion", "27-abr-2026", "Compromiso inmediato"],
            ["Entrega Sebastian y David", "04-may-2026", "Habilita QA"],
            ["Inicio QA", "04-may-2026 o siguiente dia habil", "Sujeto a cierre de desarrollo"],
            ["Fecha compromiso del proyecto", "20-may-2026", "Baseline formal actual"],
        ],
        [5.0, 2.5, 4.4],
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bullets_slide(
        s,
        "Dependencias y Riesgos",
        [
            "Dependencias abiertas: consultarCotizacionMoneda, contrato de orquestadorDeTransacciones y detalle del CC4 de CrearPrestamoORQ V4.",
            "La solicitud formal de la mediacion utilizarCupos debe gestionarla el frente banca con Interoperabilidad.",
            "Riesgo de reproceso si el CC4 cambia contratos o validaciones del flujo.",
            "Riesgo de retraso si las dependencias del cliente no se liberan a tiempo.",
        ],
        note="Feriados considerados: 01-may-2026, 18-may-2026, 08-jun-2026, 15-jun-2026 y 29-jun-2026.",
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bullets_slide(
        s,
        "Acuerdos y Proximos Pasos",
        [
            "Entregar primero utilizarCuposOrq y sin mediacion.",
            "Mantener el foco en las historias funcionales HU-BCO-004 a HU-BCO-015.",
            "Solicitar y obtener las dependencias externas pendientes del cliente e Interoperabilidad.",
            "Preparar la entrada de QA inmediatamente despues del cierre del bloque funcional comprometido.",
        ],
    )

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
