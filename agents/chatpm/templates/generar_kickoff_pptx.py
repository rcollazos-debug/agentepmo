#!/usr/bin/env python3
"""
Generador dinámico de presentación Kickoff — VortexBird PMO
Lee datos del proyecto activo desde los archivos de contexto.
Requiere: pip install python-pptx

Uso:
  python3 generar_kickoff_pptx.py                      # usa active-project.md
  python3 generar_kickoff_pptx.py --project ciudad-uao  # proyecto específico
"""

import os
import re
import argparse
from dataclasses import dataclass, field
from datetime import date
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ─── Rutas base ───────────────────────────────────────────────────────────────
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CHATPM_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))       # directorio de trabajo de ChatPM
PROJECT_ROOT = os.path.abspath(os.path.join(SCRIPT_DIR, "../../..")) # PlantillaAgente/


@dataclass
class ProjectData:
    project_id: str = "proyecto"
    project_name: str = "Proyecto"
    client: str = "Cliente"
    pm: str = "Project Manager"
    sponsor: str = "Sponsor"
    tech_lead: str = "Tech Lead"
    start_date: str = str(date.today())
    end_date: str = ""
    budget: str = "$0 USD"
    methodology: str = "Scrum / Kanban"
    technology: str = "Por definir"
    description: str = "Plataforma de software empresarial"
    team: list = field(default_factory=list)
    milestones: list = field(default_factory=list)
    risks: list = field(default_factory=list)


def _read_file(path: str) -> str:
    """Lee un archivo de texto; retorna vacío si no existe."""
    try:
        with open(path, encoding="utf-8") as f:
            return f.read()
    except FileNotFoundError:
        return ""


def _extract_field(text: str, key: str, default: str = "") -> str:
    """Extrae valor de una línea 'key: valor' en markdown."""
    pattern = rf"^\s*[-*]?\s*\*{{0,2}}{re.escape(key)}\*{{0,2}}\s*[:\|]\s*(.+)"
    for line in text.splitlines():
        m = re.match(pattern, line, re.IGNORECASE)
        if m:
            return m.group(1).strip().strip("*").strip()
    return default


def read_project_context(project_id: str = None) -> ProjectData:
    """Lee los archivos de contexto y retorna un ProjectData poblado."""
    # Resolver project_id desde active-project.md si no se especificó
    if not project_id:
        active = _read_file(os.path.join(CHATPM_DIR, "active-project.md"))
        project_id = _extract_field(active, "project_id", "proyecto")

    project_dir = os.path.join(CHATPM_DIR, "projects", project_id)
    context_dir = os.path.join(project_dir, "context")

    base = _read_file(os.path.join(context_dir, "proyecto-base.md"))
    stakeholders = _read_file(os.path.join(context_dir, "stakeholders.md"))
    dev_team = _read_file(os.path.join(context_dir, "equipodetrabajodev.md"))
    risks_file = _read_file(os.path.join(project_dir, "risks", "risk-register.md"))
    cronograma = _read_file(os.path.join(project_dir, "data", "cronograma.md"))

    data = ProjectData()
    data.project_id = project_id

    # Campos desde proyecto-base.md
    data.project_name = _extract_field(base, "Nombre del Proyecto",
                        _extract_field(base, "Nombre", project_id))
    data.client = _extract_field(base, "Cliente", "Cliente")
    data.pm = _extract_field(base, "PM",
              _extract_field(base, "Project Manager", "PM"))
    data.start_date = _extract_field(base, "Inicio", str(date.today()))
    data.end_date = _extract_field(base, "Go-Live",
                   _extract_field(base, "Fin", "Por definir"))
    data.budget = _extract_field(base, "Presupuesto",
                  _extract_field(base, "Budget", "$0 USD"))
    data.methodology = _extract_field(base, "Metodología", "Scrum / Kanban")
    data.technology = _extract_field(base, "Tecnología",
                      _extract_field(base, "Tech Stack", "Por definir"))
    data.description = _extract_field(base, "Descripción",
                       f"Plataforma de software para {data.client}")

    # Campos desde stakeholders.md
    data.sponsor = _extract_field(stakeholders, "Sponsor", "Por definir")
    data.tech_lead = _extract_field(stakeholders, "Tech Lead", "Por definir")

    # Equipo desde dev_team o stakeholders
    team_members = []
    for line in (dev_team or stakeholders).splitlines():
        if re.match(r"^\s*[-*|]", line) and "|" in line:
            cols = [c.strip() for c in line.strip("|").split("|")]
            if len(cols) >= 2 and cols[0] and cols[0] not in ("Nombre", "---", "ID"):
                team_members.append(f"{cols[0]} — {cols[1]}")
    data.team = team_members[:6] if team_members else [
        f"{data.pm} — Project Manager",
        f"{data.tech_lead} — Tech Lead",
        "Equipo de desarrollo — Por confirmar",
    ]

    # Hitos desde cronograma.md
    milestones = []
    for line in cronograma.splitlines():
        if "|" in line and ("🎯" in line or "M1" in line or "M2" in line
                            or "Go-live" in line.lower() or "MVP" in line):
            cols = [c.strip() for c in line.strip("|").split("|")]
            if len(cols) >= 3 and cols[0] and cols[0] not in ("Hito", "#", "---"):
                milestones.append((cols[0], cols[1] if len(cols) > 1 else ""))
    data.milestones = milestones[:5] if milestones else [
        ("Acta firmada", data.start_date),
        ("Arquitectura aprobada", ""),
        ("MVP funcional", ""),
        ("Inicio UAT", ""),
        ("Go-live 🎯", data.end_date),
    ]

    # Riesgos top desde risk-register.md
    risks = []
    for line in risks_file.splitlines():
        if "|" in line and ("Alto" in line or "Crítico" in line or "Medio" in line):
            cols = [c.strip() for c in line.strip("|").split("|")]
            if len(cols) >= 4 and cols[0].startswith("R-"):
                nivel = "⚠️ ALTO" if "Alto" in line or "Crítico" in line else "⚠️ MEDIO"
                risks.append((cols[0], nivel, cols[2] if len(cols) > 2 else ""))
    data.risks = risks[:3] if risks else [
        ("R-001", "⚠️ ALTO", "Retraso en entrega de requisitos del cliente"),
        ("R-002", "⚠️ ALTO", "Cambios de alcance sin control formal"),
        ("R-003", "⚠️ MEDIO", "Disponibilidad del equipo"),
    ]

    return data


# ─── VortexBird Brand System ──────────────────────────────────────────────────
PRIMARY     = RGBColor(0x10, 0xC1, 0xBC)   # Teal #10C1BC
SECONDARY   = RGBColor(0xFF, 0x91, 0x41)   # Orange #FF9141
DARK_NAVY   = RGBColor(0x1E, 0x29, 0x3A)   # Dark Navy #1E293A
TEXT_GRAY   = RGBColor(0x56, 0x56, 0x56)   # Gray #565656
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG    = RGBColor(0xF2, 0xF2, 0xF2)
TEAL_LIGHT  = RGBColor(0xE0, 0xF7, 0xF7)
RISK_RED    = RGBColor(0xC0, 0x00, 0x00)
RISK_ORANGE = RGBColor(0xFF, 0x8C, 0x00)
RISK_BG_RED    = RGBColor(0xFF, 0xE8, 0xE8)
RISK_BG_ORANGE = RGBColor(0xFF, 0xF0, 0xD8)
FONT_TITLE  = "Red Hat Display"
FONT_BODY   = "Open Sans"

LOGO_PATH  = os.path.join(CHATPM_DIR, "assets", "logo-vortexbird.png")
ICONS_DIR  = os.path.join(CHATPM_DIR, "assets", "icons")

SLIDE_W  = Inches(13.333)
SLIDE_H  = Inches(7.5)
HEADER_H = Inches(1.1)
ACCENT_H = Inches(0.07)

_ICON_NAMES = {
    'purpose':    'target',
    'scope':      'compass',
    'team':       'conference-call',
    'schedule':   'calendar',
    'budget':     'money',
    'risks':      'warning-shield',
    'success':    'checkmark--v1',
    'next_steps': 'circled-right-2',
    'contacts':   'contact-card',
}


def _get_icon_path(key, color="FFFFFF"):
    name = _ICON_NAMES.get(key)
    if not name:
        return None
    dest = os.path.join(ICONS_DIR, f"{key}.png")
    if os.path.exists(dest) and os.path.getsize(dest) > 200:
        return dest
    os.makedirs(ICONS_DIR, exist_ok=True)
    url = f"https://img.icons8.com/ios-filled/64/{color}/{name}.png"
    try:
        import urllib.request
        urllib.request.urlretrieve(url, dest)
        if os.path.getsize(dest) > 200:
            return dest
        os.remove(dest)
    except Exception:
        pass
    return None


def _add_rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    return shape


def add_header_bar(slide, title_text, slide_type=None):
    """Full Dark Navy header: icon (left) · title · logo (right) + teal accent bottom."""
    _add_rect(slide, Inches(0), Inches(0), SLIDE_W, HEADER_H, DARK_NAVY)
    _add_rect(slide, Inches(0), HEADER_H - Inches(0.045), SLIDE_W, Inches(0.045), PRIMARY)

    icon_size = Inches(0.58)
    icon_x = Inches(0.28)
    icon_y = (HEADER_H - icon_size) / 2 - Inches(0.02)
    icon_path = _get_icon_path(slide_type) if slide_type else None
    if icon_path:
        slide.shapes.add_picture(icon_path, icon_x, icon_y, width=icon_size, height=icon_size)

    title_x = Inches(1.05) if icon_path else Inches(0.4)
    title_w = Inches(10.3) if icon_path else Inches(11.0)
    tb = slide.shapes.add_textbox(title_x, Inches(0.15), title_w, Inches(0.85))
    p = tb.text_frame.paragraphs[0]
    p.text = title_text; p.font.size = Pt(28); p.font.bold = True
    p.font.color.rgb = WHITE; p.font.name = FONT_TITLE

    logo_w = Inches(1.55)
    logo_h = logo_w * (300 / 1200)
    logo_y = (HEADER_H - logo_h) / 2
    add_logo(slide, x=Inches(11.72), y=logo_y, w=logo_w)


def add_logo(slide, x=Inches(11.72), y=Inches(0.36), w=Inches(1.55)):
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, x, y, width=w)


def add_bottom_accent(slide):
    _add_rect(slide, Inches(0), SLIDE_H - ACCENT_H, SLIDE_W, ACCENT_H, SECONDARY)


_TITLE_TO_TYPE = {
    "AGENDA": None,
    "PRESENTACIÓN DEL EQUIPO": 'team',
    "DESCRIPCIÓN DEL PROYECTO": 'purpose',
    "CRONOGRAMA Y HITOS": 'schedule',
    "METODOLOGÍA": None,
    "RIESGOS IDENTIFICADOS": 'risks',
    "PRÓXIMOS PASOS": 'next_steps',
}

def add_title(slide, text, color=None):
    """Legacy helper — delegates to branded header bar with auto icon lookup."""
    add_header_bar(slide, text, _TITLE_TO_TYPE.get(text))
    add_bottom_accent(slide)


def add_person_card(slide, x, y, name, role, desc):
    _add_rect(slide, Inches(x), Inches(y), Inches(2.8), Inches(2.2), LIGHT_BG)
    _add_rect(slide, Inches(x), Inches(y), Inches(0.08), Inches(2.2), PRIMARY)
    tb = slide.shapes.add_textbox(Inches(x+0.18), Inches(y+0.12), Inches(2.55), Inches(2.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = name; p.font.size = Pt(14); p.font.bold = True
    p.font.color.rgb = DARK_NAVY; p.font.name = FONT_TITLE
    p = tf.add_paragraph(); p.text = role; p.font.size = Pt(12); p.font.bold = True
    p.font.color.rgb = PRIMARY; p.font.name = FONT_BODY
    p = tf.add_paragraph(); p.text = desc; p.font.size = Pt(11)
    p.font.color.rgb = TEXT_GRAY; p.font.name = FONT_BODY


def add_team_box(slide, x, y, title, members):
    _add_rect(slide, Inches(x), Inches(y), Inches(4.8), Inches(2.5), TEAL_LIGHT)
    _add_rect(slide, Inches(x), Inches(y), Inches(0.08), Inches(2.5), PRIMARY)
    tb = slide.shapes.add_textbox(Inches(x+0.22), Inches(y+0.15), Inches(4.4), Inches(2.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(14); p.font.bold = True
    p.font.color.rgb = PRIMARY; p.font.name = FONT_TITLE
    for m in members:
        p = tf.add_paragraph(); p.text = f"• {m}"; p.font.size = Pt(12)
        p.font.color.rgb = DARK_NAVY; p.font.name = FONT_BODY


def add_risk_card(slide, x, y, rid, nivel, desc):
    is_high = "ALTO" in nivel or "CRÍTICO" in nivel
    bg_c = RISK_BG_RED if is_high else RISK_BG_ORANGE
    border_c = RISK_RED if is_high else RISK_ORANGE
    _add_rect(slide, Inches(x), Inches(y), Inches(12.5), Inches(1.4), bg_c)
    _add_rect(slide, Inches(x), Inches(y), Inches(0.1), Inches(1.4), border_c)
    tb = slide.shapes.add_textbox(Inches(x+0.22), Inches(y+0.1), Inches(2.8), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]; p.text = rid; p.font.size = Pt(16); p.font.bold = True
    p.font.color.rgb = border_c; p.font.name = FONT_TITLE
    p = tf.add_paragraph(); p.text = nivel; p.font.size = Pt(13); p.font.bold = True
    p.font.color.rgb = border_c; p.font.name = FONT_BODY
    tb2 = slide.shapes.add_textbox(Inches(x+3.2), Inches(y+0.2), Inches(9.0), Inches(1.0))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p = tf2.paragraphs[0]; p.text = desc; p.font.size = Pt(14)
    p.font.color.rgb = DARK_NAVY; p.font.name = FONT_BODY


# ─── Generación de presentación ───────────────────────────────────────────────
def create_presentation(data: ProjectData) -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ── Slide 1: Portada ──────────────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # Full dark background
    bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = DARK_NAVY
    # Top teal bar (thin)
    _add_rect(s, Inches(0), Inches(0), SLIDE_W, Inches(0.08), PRIMARY)
    # Logo centered top
    add_logo(s, x=Inches(5.67), y=Inches(0.3), w=Inches(2.0))
    # Teal divider
    _add_rect(s, Inches(1.0), Inches(1.45), Inches(11.333), Inches(0.06), PRIMARY)
    # Project name
    tb = s.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(12.3), Inches(1.6))
    p = tb.text_frame.paragraphs[0]; p.text = data.project_name
    p.font.size = Pt(54); p.font.bold = True; p.font.color.rgb = WHITE
    p.font.name = FONT_TITLE; p.alignment = PP_ALIGN.CENTER
    # Tagline / start date
    tb3 = s.shapes.add_textbox(Inches(0.5), Inches(3.4), Inches(12.3), Inches(0.7))
    p = tb3.text_frame.paragraphs[0]; p.text = f"KICK-OFF  ·  {data.start_date}"
    p.font.size = Pt(22); p.font.color.rgb = PRIMARY; p.font.name = FONT_TITLE
    p.alignment = PP_ALIGN.CENTER
    # Meta info
    tb4 = s.shapes.add_textbox(Inches(0.5), Inches(4.4), Inches(12.3), Inches(1.5))
    tf = tb4.text_frame
    for i, (label, val) in enumerate([
        (f"PM: {data.pm}", f"Sponsor: {data.sponsor}"),
        (f"Cliente: {data.client}", f"Go-Live: {data.end_date}"),
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"  {label}     |     {val}"
        p.font.size = Pt(17); p.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        p.font.name = FONT_BODY; p.alignment = PP_ALIGN.CENTER
    # Bottom orange bar
    _add_rect(s, Inches(0), SLIDE_H - Inches(0.08), SLIDE_W, Inches(0.08), SECONDARY)

    # ── Slide 2: Agenda ───────────────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "AGENDA")
    agenda = ["1. Bienvenida y objetivos","2. Presentación del equipo",
              "3. Descripción del proyecto","4. Alcance y restricciones",
              "5. Cronograma y hitos","6. Metodología de trabajo",
              "7. Comunicación y gobernanza","8. Riesgos identificados",
              "9. Próximos pasos","10. Preguntas y discusión"]
    tb = s.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6), Inches(5))
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(agenda):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.font.size = Pt(18); p.space_after = Pt(6)
    tb2 = s.shapes.add_textbox(Inches(8.5), Inches(3), Inches(4), Inches(1.5))
    tf2 = tb2.text_frame
    p = tf2.paragraphs[0]; p.text = "Duración"; p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = PRIMARY
    p.font.name = FONT_TITLE
    p = tf2.add_paragraph(); p.text = "80 minutos"; p.font.size = Pt(36); p.font.bold = True
    p.font.color.rgb = DARK_NAVY; p.font.name = FONT_TITLE

    # ── Slide 3: Equipo ───────────────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "PRESENTACIÓN DEL EQUIPO")
    add_person_card(s, 0.5, 1.8, data.sponsor, "Sponsor", "Autoridad ejecutiva máxima")
    add_person_card(s, 3.5, 1.8, data.pm, "Project Manager", "Gestión y coordinación")
    add_person_card(s, 6.5, 1.8, data.tech_lead, "Tech Lead", "Arquitectura y decisiones técnicas")
    add_person_card(s, 9.5, 1.8, data.client, "Cliente / PO", "Requisitos y validación")
    add_team_box(s, 0.5, 4.4, "EQUIPO DE DESARROLLO", data.team)

    # ── Slide 4: Descripción ──────────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "DESCRIPCIÓN DEL PROYECTO")
    tb = s.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(12), Inches(1.5))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = f"¿Qué es {data.project_name}?"
    p.font.size = Pt(24); p.font.bold = True; p.font.color.rgb = PRIMARY; p.font.name = FONT_TITLE
    p = tf.add_paragraph(); p.text = data.description; p.font.size = Pt(20)
    p.font.color.rgb = DARK_NAVY; p.font.name = FONT_BODY
    tb2 = s.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(12), Inches(0.5))
    p = tb2.text_frame.paragraphs[0]; p.text = "DATOS CLAVE"
    p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = PRIMARY; p.font.name = FONT_TITLE
    tb3 = s.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(12), Inches(2.8))
    tf3 = tb3.text_frame; tf3.word_wrap = True
    for label, val in [("Tecnología", data.technology), ("Metodología", data.methodology),
                       ("Presupuesto", data.budget), ("Go-live", data.end_date)]:
        p = tf3.add_paragraph(); p.text = f"• {label}: {val}"; p.font.size = Pt(18)

    # ── Slide 5: Cronograma ───────────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "CRONOGRAMA Y HITOS")
    _add_rect(s, Inches(0.5), Inches(2.5), Inches(12.0), Inches(0.06), PRIMARY)
    n = len(data.milestones)
    for i, (nombre, fecha) in enumerate(data.milestones):
        x = 0.5 + (11 / max(n-1, 1)) * i
        c = s.shapes.add_shape(9, Inches(x), Inches(2.35), Inches(0.4), Inches(0.4))
        c.fill.solid(); c.fill.fore_color.rgb = PRIMARY; c.line.fill.background()
        tb = s.shapes.add_textbox(Inches(x-0.5), Inches(2.85), Inches(1.5), Inches(0.5))
        p = tb.text_frame.paragraphs[0]; p.text = fecha; p.font.size = Pt(11)
        p.font.bold = True; p.font.color.rgb = SECONDARY; p.font.name = FONT_BODY
        p.alignment = PP_ALIGN.CENTER
        tb2 = s.shapes.add_textbox(Inches(x-0.7), Inches(3.4), Inches(2), Inches(0.8))
        tb2.text_frame.word_wrap = True
        p = tb2.text_frame.paragraphs[0]; p.text = nombre; p.font.size = Pt(11)
        p.font.color.rgb = DARK_NAVY; p.font.name = FONT_BODY; p.alignment = PP_ALIGN.CENTER

    # ── Slide 6: Metodología ──────────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "METODOLOGÍA")
    _add_rect(s, Inches(0.5), Inches(1.8), Inches(5.9), Inches(2.5), TEAL_LIGHT)
    _add_rect(s, Inches(0.5), Inches(1.8), Inches(0.08), Inches(2.5), PRIMARY)
    tb = s.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(5.4), Inches(2.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]; p.text = "SCRUM"; p.font.size = Pt(24); p.font.bold = True
    p.font.color.rgb = PRIMARY; p.font.name = FONT_TITLE
    for item in ["Sprints de 2 semanas","Sprint Planning","Sprint Review","Daily Standup","Retrospectiva"]:
        p = tf.add_paragraph(); p.text = f"• {item}"; p.font.size = Pt(15)
        p.font.color.rgb = DARK_NAVY; p.font.name = FONT_BODY
    _add_rect(s, Inches(6.9), Inches(1.8), Inches(5.9), Inches(2.5), RGBColor(0xFF, 0xF3, 0xE8))
    _add_rect(s, Inches(6.9), Inches(1.8), Inches(0.08), Inches(2.5), SECONDARY)
    tb2 = s.shapes.add_textbox(Inches(7.15), Inches(2.0), Inches(5.4), Inches(2.2))
    tf2 = tb2.text_frame
    p = tf2.paragraphs[0]; p.text = "KANBAN"; p.font.size = Pt(24); p.font.bold = True
    p.font.color.rgb = SECONDARY; p.font.name = FONT_TITLE
    for item in ["Flujo continuo","WIP Limits","Tableros visuales","CFD — detección cuellos de botella"]:
        p = tf2.add_paragraph(); p.text = f"• {item}"; p.font.size = Pt(15)
        p.font.color.rgb = DARK_NAVY; p.font.name = FONT_BODY

    # ── Slide 7: Riesgos ─────────────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "RIESGOS IDENTIFICADOS")
    y = 1.8
    for rid, nivel, desc in data.risks:
        add_risk_card(s, 0.5, y, rid, nivel, desc)
        y += 1.65

    # ── Slide 8: Próximos pasos ───────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "PRÓXIMOS PASOS")
    pasos = [("1", "Confirmar herramientas de gestión", data.tech_lead),
             ("2", "Firmar acta de constitución", data.pm),
             ("3", "Definir horarios de standup", data.pm),
             ("4 ⚠️", "Entregar requisitos al equipo", data.client)]
    y = 1.35
    for num, accion, resp in pasos:
        is_urgent = "⚠️" in num
        bg_c = RISK_BG_ORANGE if is_urgent else LIGHT_BG
        border_c = SECONDARY if is_urgent else PRIMARY
        _add_rect(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.78), bg_c)
        _add_rect(s, Inches(0.5), Inches(y), Inches(0.1), Inches(0.78), border_c)
        num_box = s.shapes.add_textbox(Inches(0.7), Inches(y+0.08), Inches(0.5), Inches(0.55))
        p = num_box.text_frame.paragraphs[0]; p.text = num.replace("⚠️","").strip()
        p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = border_c; p.font.name = FONT_TITLE
        tb = s.shapes.add_textbox(Inches(1.35), Inches(y+0.06), Inches(11.3), Inches(0.65))
        tf = tb.text_frame
        p = tf.paragraphs[0]; p.text = accion; p.font.size = Pt(15); p.font.bold = True
        p.font.color.rgb = DARK_NAVY; p.font.name = FONT_BODY
        p = tf.add_paragraph(); p.text = f"→ Responsable: {resp}"; p.font.size = Pt(12)
        p.font.color.rgb = TEXT_GRAY; p.font.name = FONT_BODY
        y += 1.0

    # ── Slide 9: Q&A ─────────────────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = DARK_NAVY
    _add_rect(s, Inches(0), Inches(0), SLIDE_W, Inches(0.08), PRIMARY)
    add_logo(s, x=Inches(5.67), y=Inches(0.5), w=Inches(2.0))
    _add_rect(s, Inches(2.0), Inches(1.75), Inches(9.333), Inches(0.06), PRIMARY)
    tb = s.shapes.add_textbox(Inches(0), Inches(2.1), SLIDE_W, Inches(1.5))
    p = tb.text_frame.paragraphs[0]; p.text = "Preguntas y Respuestas"
    p.font.size = Pt(52); p.font.bold = True; p.font.color.rgb = WHITE
    p.font.name = FONT_TITLE; p.alignment = PP_ALIGN.CENTER
    tb2 = s.shapes.add_textbox(Inches(0), Inches(3.9), SLIDE_W, Inches(0.8))
    p = tb2.text_frame.paragraphs[0]; p.text = "Construimos tecnología. Construimos relaciones."
    p.font.size = Pt(22); p.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    p.font.name = FONT_BODY; p.alignment = PP_ALIGN.CENTER
    _add_rect(s, Inches(0), SLIDE_H - Inches(0.08), SLIDE_W, Inches(0.08), SECONDARY)

    # ─── Guardar ──────────────────────────────────────────────────────────────
    output_dir = os.path.join(PROJECT_ROOT, "output")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, f"{data.project_id}_kickoff.pptx")
    prs.save(output_path)
    print(f"✓ Presentación guardada: {output_path}")
    return output_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Genera presentación Kickoff del proyecto")
    parser.add_argument("--project", default=None, help="ID del proyecto (ej: ciudad-uao)")
    args = parser.parse_args()

    print(f"Leyendo contexto del proyecto...")
    data = read_project_context(args.project)
    print(f"Proyecto: {data.project_name} | Cliente: {data.client} | PM: {data.pm}")
    create_presentation(data)
    print("✓ ¡Presentación creada exitosamente!")
