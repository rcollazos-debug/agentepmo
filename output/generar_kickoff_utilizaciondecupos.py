from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

AZUL = RGBColor(0, 87, 146)
AZUL_OSCURO = RGBColor(16, 44, 84)
BLANCO = RGBColor(255, 255, 255)
GRIS = RGBColor(90, 90, 90)
VERDE = RGBColor(34, 139, 34)
NARANJA = RGBColor(219, 131, 0)
ROJO = RGBColor(180, 45, 45)

data = {
    "nombre": "Utilización de Cupos",
    "cliente": "Bancoomeva",
    "pm": "Raul Collazos",
    "sponsor": "Keilly Barona",
    "po": "Ingrid Y. Mosquera",
    "fecha_inicio": "Por confirmar",
    "fecha_fin": "2.7 meses base / 2 meses solicitado por cliente",
    "presupuesto": "COP $112.722.120",
    "proposito": "Implementar un módulo para convertir cupos de crédito preaprobados en préstamos desembolsados, con operación segura, parametrizable, trazable e integrada con COBIS.",
    "objetivos": [
        "Reducir tiempos de desembolso y eliminar pasos manuales.",
        "Integrar el flujo con COBIS y servicios corporativos.",
        "Centralizar reglas de negocio con parametrización de cupos y cuentas.",
        "Entregar solución validada en ambiente DEV con evidencias de pruebas."
    ],
    "incluye": [
        "MFE Front para identificación, consulta y selección de cupos.",
        "Back MFE con validaciones, reglas de negocio y auditoría.",
        "Parametrización CRUD de cupos y cuentas.",
        "Desembolso por cuenta, cheque y mixto.",
        "Validaciones de cliente, mora y embargo."
    ],
    "excluye": [
        "Mediación utilizarCupos.",
        "Mediación para consulta de parametrizaciones.",
        "Desarrollos en legados/core fuera de alcance.",
        "Licencias, hardware e infraestructura cloud."
    ],
    "equipo": [
        ("Raul Collazos", "Project Manager"),
        ("Ingrid Y. Mosquera", "Product Owner"),
        ("Sebastian Garcia", "Tech Lead VortexBird"),
        ("Adriana Muñoz Ñañez", "Ingeniería Soluciones TI"),
        ("Monica Araque Urrego", "STE Banca Personas Activas"),
        ("Por definir", "Equipo Dev / QA / UX")
    ],
    "hitos": [
        ("Requerimientos base", "18-mar-2026"),
        ("Ajustes principales HUs", "09-abr-2026"),
        ("Propuesta presentada", "16-abr-2026"),
        ("Inicio proyecto", "Por confirmar"),
        ("MVP funcional", "Por definir"),
        ("Entrega cliente solicitada", "2 meses"),
        ("Entrega DEV estimada", "+2.7 meses")
    ],
    "riesgos": [
        ("Dependencia de Interoperabilidad", "Alto", "Acordar responsable, fechas y seguimiento semanal."),
        ("Requisitos incompletos", "Alto", "Refinamiento temprano y cierre funcional por sprint."),
        ("Insumos técnicos incompletos", "Alto", "Checklist de contratos/servicios y escalación con cliente."),
        ("Bloqueos por integraciones", "Alto", "Validación temprana de contratos, mocks y pruebas DEV."),
        ("Compresión de cronograma", "Alto", "Evaluar MVP, capacidad adicional y formalizar cambio si aplica.")
    ],
    "criterios": [
        "Flujo end-to-end ejecutable en DEV.",
        "Validaciones de cliente, mora y embargo operativas.",
        "Integraciones críticas funcionando según alcance.",
        "Pruebas unitarias, integración y funcionales con evidencia.",
        "Trazabilidad técnica y operativa disponible."
    ],
    "proximos": [
        "Confirmar responsable y calendario de Interoperabilidad.",
        "Definir si la meta de 2 meses es expectativa o cambio formal.",
        "Cerrar pendientes funcionales de HUs y reglas de campo.",
        "Asegurar accesos, ambientes y credenciales desde semana 1.",
        "Definir backlog inicial, MVP y plan de sprints.",
        "Establecer control de cambios formal."
    ]
}

def textbox(slide, x, y, w, h, text, size=18, bold=False, color=GRIS, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return tf

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 1 portada
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(7.5))
bg.fill.solid(); bg.fill.fore_color.rgb = AZUL_OSCURO; bg.line.fill.background()
textbox(s, 0.6, 1.1, 12, 1, "KICKOFF MEETING", 28, True, BLANCO, PP_ALIGN.CENTER)
textbox(s, 0.6, 2.1, 12, 1.2, data["nombre"], 30, True, BLANCO, PP_ALIGN.CENTER)
textbox(s, 0.6, 3.5, 12, 1.2, f"Cliente: {data['cliente']}\nPM: {data['pm']} | PO: {data['po']}\nSponsor: {data['sponsor']}", 18, False, BLANCO, PP_ALIGN.CENTER)
textbox(s, 0.6, 5.5, 12, 0.6, "Borrador PMO — 19-abr-2026", 16, False, BLANCO, PP_ALIGN.CENTER)

# 2 propósito y objetivos
s = prs.slides.add_slide(prs.slide_layouts[6])
textbox(s, 0.5, 0.3, 12, 0.6, "Propósito y Objetivos", 28, True, AZUL_OSCURO)
textbox(s, 0.6, 1.2, 12, 1.2, data["proposito"], 20)
tf = textbox(s, 0.8, 2.8, 12, 3.8, "", 18)
for i, item in enumerate(data["objetivos"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = f"• {item}"
    p.font.size = Pt(18)

# 3 alcance
s = prs.slides.add_slide(prs.slide_layouts[6])
textbox(s, 0.5, 0.3, 12, 0.6, "Alcance y Exclusiones", 28, True, AZUL_OSCURO)
box1 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(6), Inches(5.7))
box1.fill.solid(); box1.fill.fore_color.rgb = RGBColor(232, 245, 233); box1.line.color.rgb = VERDE
box2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.2), Inches(6), Inches(5.7))
box2.fill.solid(); box2.fill.fore_color.rgb = RGBColor(255, 243, 224); box2.line.color.rgb = NARANJA
textbox(s, 0.8, 1.4, 3, 0.4, "Incluye", 20, True, VERDE)
tf = textbox(s, 0.8, 1.9, 5.3, 4.8, "", 16)
for i, item in enumerate(data["incluye"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = f"• {item}"; p.font.size = Pt(16)
textbox(s, 7.1, 1.4, 3, 0.4, "Excluye", 20, True, NARANJA)
tf = textbox(s, 7.1, 1.9, 5.2, 4.8, "", 16)
for i, item in enumerate(data["excluye"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = f"• {item}"; p.font.size = Pt(16)

# 4 equipo
s = prs.slides.add_slide(prs.slide_layouts[6])
textbox(s, 0.5, 0.3, 12, 0.6, "Equipo y Stakeholders Clave", 28, True, AZUL_OSCURO)
y = 1.2
for idx, (name, role) in enumerate(data["equipo"]):
    x = 0.6 + (idx % 2) * 6.2
    if idx and idx % 2 == 0:
        y += 1.5
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.8), Inches(1.2))
    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(240, 248, 255); shp.line.color.rgb = AZUL
    textbox(s, x + 0.15, y + 0.1, 5.4, 0.4, name, 18, True, AZUL_OSCURO)
    textbox(s, x + 0.15, y + 0.5, 5.4, 0.35, role, 14, False, AZUL)

# 5 hitos
s = prs.slides.add_slide(prs.slide_layouts[6])
textbox(s, 0.5, 0.3, 12, 0.6, "Cronograma General e Hitos", 28, True, AZUL_OSCURO)
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.2), Inches(11.4), Inches(0.08))
line.fill.solid(); line.fill.fore_color.rgb = AZUL; line.line.fill.background()
for i, (hito, fecha) in enumerate(data["hitos"]):
    x = 0.9 + i * 2.15
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(3.0), Inches(0.28), Inches(0.28))
    dot.fill.solid(); dot.fill.fore_color.rgb = AZUL_OSCURO; dot.line.fill.background()
    textbox(s, x - 0.4, 2.2, 1.1, 0.5, fecha, 11, True, AZUL_OSCURO, PP_ALIGN.CENTER)
    textbox(s, x - 0.65, 3.45, 1.6, 0.9, hito, 11, False, GRIS, PP_ALIGN.CENTER)
textbox(s, 0.8, 5.55, 12, 0.8, f"Plan base: 2.7 meses | Solicitud cliente: 2 meses\nDuración de referencia: {data['fecha_fin']}", 17, True, AZUL)

# 6 presupuesto y tecnología
s = prs.slides.add_slide(prs.slide_layouts[6])
textbox(s, 0.5, 0.3, 12, 0.6, "Presupuesto y Stack Tecnológico", 28, True, AZUL_OSCURO)
shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.2), Inches(4.2), Inches(2))
shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(230, 245, 255); shp.line.color.rgb = AZUL
textbox(s, 1.0, 1.5, 3.6, 0.5, "Presupuesto estimado", 20, True, AZUL)
textbox(s, 1.0, 2.1, 3.6, 0.6, data["presupuesto"], 24, True, AZUL_OSCURO)
tf = textbox(s, 5.4, 1.3, 7, 4.8, "", 17)
for i, item in enumerate([
    "Frontend: Angular, TypeScript",
    "Backend: Java 25, Spring Boot, Spring Security, JPA",
    "Base de datos: PostgreSQL",
    "DevOps: Git, Jenkins, Maven",
    "Observabilidad: Grafana/ELK",
    "Seguridad: Vault",
    "Arquitectura: Microservicios / contenedorizable"
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = f"• {item}"; p.font.size = Pt(17)

# 7 riesgos
s = prs.slides.add_slide(prs.slide_layouts[6])
textbox(s, 0.5, 0.3, 12, 0.6, "Top Riesgos y Mitigación", 28, True, AZUL_OSCURO)
y = 1.1
for nombre, nivel, mitigacion in data["riesgos"]:
    color = RGBColor(255, 235, 238) if nivel == "Alto" else RGBColor(255, 248, 225)
    borde = ROJO if nivel == "Alto" else NARANJA
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(y), Inches(12.1), Inches(1.0))
    shp.fill.solid(); shp.fill.fore_color.rgb = color; shp.line.color.rgb = borde
    textbox(s, 0.8, y + 0.12, 4.2, 0.3, nombre, 16, True, AZUL_OSCURO)
    textbox(s, 4.6, y + 0.12, 1.2, 0.3, nivel, 14, True, borde)
    textbox(s, 5.7, y + 0.12, 6.6, 0.5, mitigacion, 14, False, GRIS)
    y += 1.1

# 8 criterios de éxito
s = prs.slides.add_slide(prs.slide_layouts[6])
textbox(s, 0.5, 0.3, 12, 0.6, "Criterios de Éxito", 28, True, AZUL_OSCURO)
tf = textbox(s, 0.9, 1.2, 11.8, 5.5, "", 18)
for i, item in enumerate(data["criterios"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = f"✓ {item}"; p.font.size = Pt(18)

# 9 próximos pasos
s = prs.slides.add_slide(prs.slide_layouts[6])
textbox(s, 0.5, 0.3, 12, 0.6, "Próximos Pasos Inmediatos", 28, True, AZUL_OSCURO)
y = 1.2
for i, item in enumerate(data["proximos"], start=1):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y), Inches(11.6), Inches(0.75))
    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(240, 248, 255); shp.line.color.rgb = AZUL
    textbox(s, 1.0, y + 0.12, 11, 0.3, f"{i}. {item}", 17, False, GRIS)
    y += 0.95

# 10 cierre
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(7.5))
bg.fill.solid(); bg.fill.fore_color.rgb = AZUL_OSCURO; bg.line.fill.background()
textbox(s, 0.5, 2.0, 12.3, 0.8, "Preguntas y Contactos", 30, True, BLANCO, PP_ALIGN.CENTER)
textbox(s, 0.5, 3.1, 12.3, 1.2, f"PM: {data['pm']}\nPO: {data['po']}\nCliente: {data['cliente']}", 18, False, BLANCO, PP_ALIGN.CENTER)
textbox(s, 0.5, 5.2, 12.3, 0.6, "VortexBird PMO — borrador kickoff", 16, False, BLANCO, PP_ALIGN.CENTER)

output = "/Users/rcollazos/Library/Mobile Documents/com~apple~CloudDocs/Documents/Proyectos/agentes/PlantillaAgente/output/utilizaciondecupos_kickoff.pptx"
prs.save(output)
print(output)
